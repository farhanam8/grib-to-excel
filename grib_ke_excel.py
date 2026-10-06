"""
GRIB ke Excel
=============
Aplikasi kecil untuk mengubah file GRIB (ERA5, ERA5-Land, dan sejenisnya)
menjadi file Excel pada satu titik lokasi.

Cara jalan:
    pip install -r requirements.txt
    streamlit run grib_ke_excel.py

Alur:
    1. Pengguna memasukkan koordinat (lintang, bujur).
    2. Aplikasi mencari titik grid terdekat untuk setiap bentuk grid di file GRIB,
       lalu membaca nilai di titik itu saja dari setiap message (via ecCodes).
    3. Variabel mentah diterjemahkan ke besaran yang mudah dibaca,
       misalnya u10 + v10 menjadi kecepatan dan arah angin.
    4. Grafik deret waktu (resolusi otomatis) dan windrose untuk data angin.
    5. Hasilnya diunduh sebagai Excel (sheet Data, Ringkasan, Windrose, Info).
"""

from __future__ import annotations

import io
import os
import tempfile
import warnings
from dataclasses import dataclass, field
from datetime import datetime

import numpy as np
import pandas as pd

warnings.filterwarnings("ignore", category=FutureWarning)

R_BUMI_KM = 6371.0088
NAMA_ARAH = ["Utara", "Timur Laut", "Timur", "Tenggara",
             "Selatan", "Barat Daya", "Barat", "Barat Laut"]
ZONA_WAKTU = {"UTC": 0, "WIB (UTC+7)": 7, "WITA (UTC+8)": 8, "WIT (UTC+9)": 9}

# --------------------------------------------------------------------------
# Kamus terjemahan variabel
# --------------------------------------------------------------------------
# Pasangan komponen angin: (u, v, keterangan ketinggian)
PASANGAN_ANGIN = [
    ("u10", "v10", "10 m"),
    ("u100", "v100", "100 m"),
    ("u10n", "v10n", "10 m netral"),
    ("u", "v", ""),          # level tekanan, label level ditambahkan otomatis
]


def _k2c(x):
    return x - 273.15


def _kali(f):
    return lambda x: x * f


# nama_pendek: (nama tampil, satuan, fungsi konversi, keterangan, jenis)
# jenis: "biasa", "jumlah" (boleh dijumlah, akumulasi), "fluks" (akumulasi
# yang dirata-ratakan), "arah" (derajat).
KONVERSI = {
    "t2m":  ("Suhu udara 2 m", "°C", _k2c, "K - 273,15", "biasa"),
    "d2m":  ("Suhu titik embun 2 m", "°C", _k2c, "K - 273,15", "biasa"),
    "skt":  ("Suhu kulit permukaan", "°C", _k2c, "K - 273,15", "biasa"),
    "sst":  ("Suhu permukaan laut", "°C", _k2c, "K - 273,15", "biasa"),
    "t":    ("Suhu udara", "°C", _k2c, "K - 273,15", "biasa"),
    "mx2t": ("Suhu maksimum 2 m", "°C", _k2c, "K - 273,15", "biasa"),
    "mn2t": ("Suhu minimum 2 m", "°C", _k2c, "K - 273,15", "biasa"),
    "stl1": ("Suhu tanah lapisan 1 (0-7 cm)", "°C", _k2c, "K - 273,15", "biasa"),
    "stl2": ("Suhu tanah lapisan 2 (7-28 cm)", "°C", _k2c, "K - 273,15", "biasa"),
    "sp":   ("Tekanan permukaan", "hPa", _kali(0.01), "Pa / 100", "biasa"),
    "msl":  ("Tekanan permukaan laut", "hPa", _kali(0.01), "Pa / 100", "biasa"),
    "tp":   ("Curah hujan total", "mm", _kali(1000), "m x 1000", "jumlah"),
    "cp":   ("Hujan konvektif", "mm", _kali(1000), "m x 1000", "jumlah"),
    "lsp":  ("Hujan skala besar", "mm", _kali(1000), "m x 1000", "jumlah"),
    "sf":   ("Salju (setara air)", "mm", _kali(1000), "m x 1000", "jumlah"),
    "ro":   ("Limpasan total", "mm", _kali(1000), "m x 1000", "jumlah"),
    "e":    ("Evaporasi (positif = menguap)", "mm", _kali(-1000), "m x -1000", "jumlah"),
    "pev":  ("Evaporasi potensial (positif = menguap)", "mm", _kali(-1000), "m x -1000", "jumlah"),
    "ssrd": ("Radiasi matahari ke permukaan", "W/m²", _kali(1 / 3600), "J/m² per jam / 3600", "fluks"),
    "strd": ("Radiasi gelombang panjang ke permukaan", "W/m²", _kali(1 / 3600), "J/m² per jam / 3600", "fluks"),
    "ssr":  ("Radiasi matahari bersih", "W/m²", _kali(1 / 3600), "J/m² per jam / 3600", "fluks"),
    "str":  ("Radiasi gelombang panjang bersih", "W/m²", _kali(1 / 3600), "J/m² per jam / 3600", "fluks"),
    "tisr": ("Radiasi matahari puncak atmosfer", "W/m²", _kali(1 / 3600), "J/m² per jam / 3600", "fluks"),
    "tcc":  ("Tutupan awan total", "%", _kali(100), "fraksi x 100", "biasa"),
    "lcc":  ("Tutupan awan rendah", "%", _kali(100), "fraksi x 100", "biasa"),
    "mcc":  ("Tutupan awan menengah", "%", _kali(100), "fraksi x 100", "biasa"),
    "hcc":  ("Tutupan awan tinggi", "%", _kali(100), "fraksi x 100", "biasa"),
    "r":    ("Kelembapan relatif", "%", None, "tanpa konversi", "biasa"),
    "q":    ("Kelembapan spesifik", "g/kg", _kali(1000), "kg/kg x 1000", "biasa"),
    "z":    ("Tinggi geopotensial", "m", _kali(1 / 9.80665), "m²/s² / 9,80665", "biasa"),
    "w":    ("Kecepatan vertikal", "Pa/s", None, "tanpa konversi", "biasa"),
    "swh":  ("Tinggi gelombang signifikan", "m", None, "tanpa konversi", "biasa"),
    "shww": ("Tinggi gelombang angin", "m", None, "tanpa konversi", "biasa"),
    "shts": ("Tinggi swell total", "m", None, "tanpa konversi", "biasa"),
    "mwp":  ("Periode gelombang rata-rata", "s", None, "tanpa konversi", "biasa"),
    "pp1d": ("Periode puncak gelombang", "s", None, "tanpa konversi", "biasa"),
    "mwd":  ("Arah gelombang rata-rata", "°", None, "derajat dari utara, arah datang", "arah"),
    "cape": ("CAPE", "J/kg", None, "tanpa konversi", "biasa"),
    "blh":  ("Tinggi lapisan batas atmosfer", "m", None, "tanpa konversi", "biasa"),
    "i10fg": ("Hembusan angin maksimum 10 m", "m/s", None, "tanpa konversi", "biasa"),
    "fg10": ("Hembusan angin maksimum 10 m", "m/s", None, "tanpa konversi", "biasa"),
    "tcwv": ("Total uap air kolom", "mm", None, "kg/m² setara mm", "biasa"),
    "swvl1": ("Kelembapan tanah lapisan 1 (0-7 cm)", "m³/m³", None, "tanpa konversi", "biasa"),
    "swvl2": ("Kelembapan tanah lapisan 2 (7-28 cm)", "m³/m³", None, "tanpa konversi", "biasa"),
}

# Variabel akumulasi yang nilainya tidak boleh negatif (sisa pembulatan GRIB)
TIDAK_NEGATIF = {"tp", "cp", "lsp", "sf", "ro", "ssrd", "strd", "tisr"}


# --------------------------------------------------------------------------
# Struktur hasil
# --------------------------------------------------------------------------
@dataclass
class InfoGrid:
    sumber: str
    variabel: str
    lat_grid: float
    lon_grid: float
    jarak_km: float
    resolusi: str
    catatan: str = ""


@dataclass
class Hasil:
    data: pd.DataFrame                      # hasil terjemahan, index waktu lokal
    mentah: pd.DataFrame                    # nilai asli, index waktu UTC
    meta: dict                              # kolom mentah -> keterangan
    satuan: dict                            # kolom hasil -> satuan
    jenis: dict                             # kolom hasil -> jenis
    konversi: list                          # (kolom, asal, keterangan)
    grid: list                              # list[InfoGrid]
    peringatan: list = field(default_factory=list)
    zona: str = "UTC"
    lat_input: float = 0.0
    lon_input: float = 0.0
    sumber: list = field(default_factory=list)
    angin: list = field(default_factory=list)   # [{"nama", "kec", "arah"}] untuk windrose


# --------------------------------------------------------------------------
# Membaca GRIB dan memilih grid
# --------------------------------------------------------------------------
# GRIB disimpan sebagai deretan "message", satu message = satu variabel pada
# satu waktu dan satu level. Aplikasi ini membaca message satu per satu lewat
# ecCodes dan hanya mengambil SATU nilai (titik grid terpilih) dari tiap message,
# tanpa membongkar seluruh peta. Ini jauh lebih cepat dan hemat memori
# dibanding membuka seluruh file dengan xarray/cfgrib.

def jarak_km(lat1, lon1, lat2, lon2):
    p1, p2 = np.radians(lat1), np.radians(lat2)
    dphi = p2 - p1
    dlam = np.radians(lon2 - lon1)
    a = np.sin(dphi / 2) ** 2 + np.cos(p1) * np.cos(p2) * np.sin(dlam / 2) ** 2
    return 2 * R_BUMI_KM * np.arcsin(np.sqrt(np.clip(a, 0, 1)))


def _ambil(ec, h, kunci, bawaan=None):
    try:
        return ec.codes_get(h, kunci) if ec.codes_is_defined(h, kunci) else bawaan
    except Exception:  # noqa: BLE001
        return bawaan


def _geometri(ec, h, lat, lon):
    """Hitung jarak input ke semua titik grid. Dipanggil sekali per bentuk grid."""
    lats = ec.codes_get_array(h, "latitudes")
    lons = ec.codes_get_array(h, "longitudes")
    d = jarak_km(lat, lon, lats, lons)
    idx0 = int(np.argmin(d))

    if _ambil(ec, h, "gridType") == "regular_ll":
        di = float(_ambil(ec, h, "iDirectionIncrementInDegrees", 0.25))
        dj = float(_ambil(ec, h, "jDirectionIncrementInDegrees", 0.25))
        resolusi, langkah = f"{di:g}° x {dj:g}°", max(di, dj)
    else:
        resolusi, langkah = str(_ambil(ec, h, "gridType", "-")), 0.5

    peringatan = ""
    if d[idx0] > 1.5 * langkah * 111.2:
        peringatan = (f"Koordinat input berada di luar cakupan data "
                      f"(lintang {lats.min():g} s.d. {lats.max():g}, bujur {lons.min():g} s.d. {lons.max():g}). "
                      f"Grid terdekat berjarak {d[idx0]:.1f} km.")
    return {"lats": lats, "lons": lons, "d": d, "idx0": idx0,
            "resolusi": resolusi, "peringatan": peringatan}


def _label_level(ec, h):
    jenis = _ambil(ec, h, "typeOfLevel", "")
    level = _ambil(ec, h, "level", 0)
    label = ""
    if jenis == "isobaricInhPa":
        label = f"{level:g} hPa"
    elif jenis in ("hybrid", "modelLevel"):
        label = f"model level {level}"
    elif jenis == "potentialVorticity":
        label = f"PV {level}"
    elif jenis == "theta":
        label = f"theta {level} K"
    anggota = _ambil(ec, h, "numberOfForecastsInEnsemble", 0) or 0
    if anggota > 1:
        no = _ambil(ec, h, "number", 0)
        label = f"{label}, anggota {no}" if label else f"anggota {no}"
    return label


def _pesan_grib(path):
    """Potong file menjadi message GRIB mentah (bytes) tanpa ecCodes.
    Menghasilkan (posisi_byte, bytes). Message GRIB1 berukuran sangat besar
    (format 'large GRIB1') dikembalikan None supaya dibaca lewat ecCodes."""
    import mmap
    with open(path, "rb") as f:
        if os.path.getsize(path) == 0:
            return
        mm = mmap.mmap(f.fileno(), 0, access=mmap.ACCESS_READ)
        try:
            pos, n = 0, len(mm)
            while True:
                pos = mm.find(b"GRIB", pos)
                if pos < 0 or pos + 16 > n:
                    break
                edisi = mm[pos + 7]
                if edisi == 1:
                    panjang = int.from_bytes(mm[pos + 4:pos + 7], "big")
                    if panjang & 0x800000:          # large GRIB1, serahkan ke ecCodes
                        yield pos, None
                        return
                elif edisi == 2:
                    panjang = int.from_bytes(mm[pos + 8:pos + 16], "big")
                else:
                    pos += 4
                    continue
                if panjang < 16 or pos + panjang > n:
                    break
                yield pos, mm[pos:pos + panjang]
                pos += panjang
        finally:
            mm.close()


def _tanda_besaran(b):
    """Bilangan bulat bertanda gaya GRIB1 (bit pertama = tanda)."""
    v = int.from_bytes(b, "big")
    bit = len(b) * 8 - 1
    return -(v & ((1 << bit) - 1)) if v >> bit else v


def _ibm_float(b):
    s = -1.0 if b[0] & 0x80 else 1.0
    eksp = (b[0] & 0x7F) - 64
    mant = int.from_bytes(b[1:4], "big")
    return s * mant * 2.0 ** -24 * 16.0 ** eksp


_SATUAN_WAKTU_GRIB1 = {0: 1 / 60, 1: 1, 2: 24, 10: 3, 11: 6, 12: 12, 13: 0.25, 14: 0.5, 254: 1 / 3600}


def _grib1_cepat(m):
    """Urai header GRIB1 dengan packing sederhana (format umum ERA5 dari CDS).
    Mengembalikan dict, atau None kalau format di luar jalur cepat."""
    pl = int.from_bytes(m[8:11], "big")
    pds = m[8:8 + pl]
    if pl < 28:
        return None
    bendera = pds[7]
    if not bendera & 0x80:                       # tanpa GDS, tidak bisa
        return None
    satuan = _SATUAN_WAKTU_GRIB1.get(pds[17])
    p1, p2, tri = pds[18], pds[19], pds[20]
    if satuan is None or tri not in (0, 1, 2, 3, 4, 5, 10):
        return None
    if tri == 10:
        geser = (p1 * 256 + p2) * satuan
    elif tri in (2, 3, 4, 5):
        geser = p2 * satuan
    elif tri == 0:
        geser = p1 * satuan
    else:
        geser = 0
    tahun = (pds[24] - 1) * 100 + pds[12]
    try:
        waktu = pd.Timestamp(tahun, pds[13], pds[14], pds[15], pds[16]) + pd.Timedelta(hours=geser)
    except ValueError:
        return None

    off = 8 + pl
    gl = int.from_bytes(m[off:off + 3], "big")
    gds = m[off:off + gl]
    off += gl
    bitmap = None
    if bendera & 0x40:
        bl = int.from_bytes(m[off:off + 3], "big")
        if int.from_bytes(m[off + 4:off + 6], "big") != 0:   # bitmap bawaan, jarang
            return None
        bitmap = m[off + 6:off + bl]
        off += bl
    if m[off + 3] & 0xF0:                        # bukan grid point simple packing
        return None
    return {
        "header": (pds[3], pds[8], pds[9], pds[10], pds[11], gds),
        "waktu": waktu,
        "langkah": float(geser),
        "D": _tanda_besaran(pds[26:28]),
        "E": _tanda_besaran(m[off + 4:off + 6]),
        "R": _ibm_float(m[off + 6:off + 10]),
        "nb": m[off + 10],
        "data": off + 11,
        "bitmap": bitmap,
    }


def _nilai_grib1(m, c, idx, cache_bitmap):
    """Ambil satu nilai dari data simple packing tanpa membongkar seluruh peta."""
    p = idx
    bm = c["bitmap"]
    if bm is not None:
        if not (bm[idx >> 3] >> (7 - (idx & 7))) & 1:
            return np.nan
        kunci = (bytes(bm), idx)
        p = cache_bitmap.get(kunci)
        if p is None:
            bit = np.unpackbits(np.frombuffer(bm, dtype=np.uint8))
            p = int(bit[:idx].sum())
            cache_bitmap[kunci] = p
    nb = c["nb"]
    x = 0
    if nb:
        posbit = p * nb
        awal = c["data"] + (posbit >> 3)
        nbyte = ((posbit & 7) + nb + 7) >> 3
        x = (int.from_bytes(m[awal:awal + nbyte], "big") >> (nbyte * 8 - (posbit & 7) - nb)) & ((1 << nb) - 1)
    return (c["R"] + x * 2.0 ** c["E"]) / 10.0 ** c["D"]


def baca_titik_grib(path, nama_tampil, lat, lon, lewati_kosong=True, kabar=None):
    """Baca deret waktu di grid terdekat dari satu file GRIB.

    Jalur cepat: message GRIB1 simple packing (format ERA5 dari CDS) diurai
    langsung dengan Python, hanya mengambil satu nilai per message. ecCodes
    cukup dipakai sekali untuk setiap kombinasi variabel/level/grid baru.
    Jalur biasa: format lain (GRIB2, packing lain) dibaca lewat ecCodes.
    kabar: fungsi opsional kabar(fraksi_0_sampai_1) untuk progress bar."""
    import eccodes as ec

    ukuran = max(os.path.getsize(path), 1)
    geo = {}            # md5 grid -> info geometri
    pilihan = {}        # (md5, variabel) -> (indeks, catatan)
    daftar_header = {}  # header GRIB1 mentah -> info variabel
    cache_bitmap = {}
    nilai, meta = {}, {}
    peringatan = set()
    jumlah_ganda = 0
    n = 0

    def info_dari_handle(h):
        """Metadata satu message lewat ecCodes, plus pemilihan grid."""
        var = _ambil(ec, h, "cfVarName", "unknown")
        if var in (None, "unknown", "~"):
            var = _ambil(ec, h, "shortName", "unknown")
        if var in (None, "unknown", "~"):
            var = f"param{_ambil(ec, h, 'paramId', n)}"
        md5 = _ambil(ec, h, "md5GridSection", "grid")
        if md5 not in geo:
            geo[md5] = _geometri(ec, h, lat, lon)
            if geo[md5]["peringatan"]:
                peringatan.add(f"{nama_tampil}: {geo[md5]['peringatan']}")
        g = geo[md5]
        ada_bitmap = bool(_ambil(ec, h, "bitmapPresent", 0))
        hilang = float(_ambil(ec, h, "missingValue", 9999))
        if (md5, var) not in pilihan:
            idx, catatan = g["idx0"], ""
            if lewati_kosong and ada_bitmap:
                isi = ec.codes_get_values(h) != hilang
                if isi.any() and not isi[idx]:
                    idx = int(np.argmin(np.where(isi, g["d"], np.inf)))
                    catatan = (f"Grid terdekat ({g['lats'][g['idx0']]:g}, {g['lons'][g['idx0']]:g}) "
                               f"tidak berisi data, dipakai grid berisi terdekat.")
            pilihan[(md5, var)] = (idx, catatan)
        label = _label_level(ec, h)
        return {
            "var": var, "md5": md5, "idx": pilihan[(md5, var)][0], "label": label,
            "kunci": f"{var}@{label}" if label else var,
            "ada_bitmap": ada_bitmap, "hilang": hilang,
            "satuan": _ambil(ec, h, "units", "") or "",
            "nama_panjang": _ambil(ec, h, "name", var) or var,
            "jenis_langkah": _ambil(ec, h, "stepType", "") or "",
        }

    def simpan(info, waktu, v, langkah):
        nonlocal jumlah_ganda
        kunci = info["kunci"]
        deret = nilai.setdefault(kunci, {})
        if waktu in deret:
            jumlah_ganda += 1
        else:
            deret[waktu] = v
        if kunci not in meta:
            meta[kunci] = {
                "nama": info["var"], "level": info["label"], "satuan": info["satuan"],
                "nama_panjang": info["nama_panjang"], "jenis_langkah": info["jenis_langkah"],
                "langkah_maks_jam": langkah, "md5": info["md5"],
            }
        elif langkah > meta[kunci]["langkah_maks_jam"]:
            meta[kunci]["langkah_maks_jam"] = langkah

    def lewat_eccodes(h):
        info = info_dari_handle(h)
        v = ec.codes_get_double_element(h, "values", info["idx"])
        if info["ada_bitmap"] and v == info["hilang"]:
            v = np.nan
        tgl = int(_ambil(ec, h, "validityDate"))
        jam = int(_ambil(ec, h, "validityTime"))
        waktu = pd.Timestamp(tgl // 10000, tgl // 100 % 100, tgl % 100, jam // 100, jam % 100)
        try:
            langkah = float(_ambil(ec, h, "endStep", 0))
        except (TypeError, ValueError):
            langkah = 0.0
        simpan(info, waktu, v, langkah)

    for pos, m in _pesan_grib(path):
        if m is None:
            # Format yang tidak bisa dipotong manual: baca sisa file lewat ecCodes
            with open(path, "rb") as f:
                f.seek(pos)
                while True:
                    h = ec.codes_grib_new_from_file(f)
                    if h is None:
                        break
                    try:
                        lewat_eccodes(h)
                    finally:
                        ec.codes_release(h)
                    n += 1
                    if kabar and n % 250 == 0:
                        kabar(min(f.tell() / ukuran, 1.0))
            break

        c = _grib1_cepat(m) if m[7] == 1 else None
        if c is not None:
            info = daftar_header.get(c["header"])
            if info is None:
                h = ec.codes_new_from_message(bytes(m))
                try:
                    info = info_dari_handle(h)
                finally:
                    ec.codes_release(h)
                daftar_header[c["header"]] = info
            simpan(info, c["waktu"], _nilai_grib1(m, c, info["idx"], cache_bitmap), c["langkah"])
        else:
            h = ec.codes_new_from_message(bytes(m))
            try:
                lewat_eccodes(h)
            finally:
                ec.codes_release(h)

        n += 1
        if kabar and n % 2000 == 0:
            kabar(min(pos / ukuran, 1.0))

    if n == 0:
        raise ValueError(f"{nama_tampil}: tidak ada message GRIB yang terbaca. Pastikan file berformat GRIB.")
    if jumlah_ganda:
        peringatan.add(f"{nama_tampil}: {jumlah_ganda} nilai ganda (variabel, level, dan waktu sama) diabaikan.")

    seri = {}
    for kunci, deret in nilai.items():
        s = pd.Series(deret, dtype=float).sort_index()
        s.index.name = "waktu"
        seri[kunci] = s

    kelompok = {}
    for (md5, var), (idx, catatan) in pilihan.items():
        kelompok.setdefault((md5, idx, catatan), []).append(var)
    grid = []
    for (md5, idx, catatan), daftar_var in kelompok.items():
        g = geo[md5]
        lon_g = float(g["lons"][idx])
        if lon_g > 180:
            lon_g -= 360
        grid.append(InfoGrid(nama_tampil, ", ".join(daftar_var), float(g["lats"][idx]), lon_g,
                             float(g["d"][idx]), g["resolusi"], catatan))
    if kabar:
        kabar(1.0)
    return seri, meta, grid, sorted(peringatan)


# --------------------------------------------------------------------------
# Terjemahan variabel
# --------------------------------------------------------------------------
def ke_mata_angin(derajat: pd.Series) -> pd.Series:
    hasil = pd.Series(pd.NA, index=derajat.index, dtype="object")
    ada = derajat.notna()
    idx = (np.floor(((derajat[ada] % 360) + 22.5) / 45).astype(int)) % 8
    hasil[ada] = [NAMA_ARAH[k] for k in idx]
    return hasil


def kumulatif_ke_per_jam(s: pd.Series) -> pd.Series:
    """ERA5-Land: akumulasi dimulai dari 00 UTC. Nilai per jam = selisih
    dengan jam sebelumnya, kecuali pukul 01 UTC yang sudah nilai 1 jam."""
    s = s.dropna().sort_index()
    if s.empty:
        return s
    per_jam = s.diff()
    jam1 = s.index.hour == 1
    per_jam[jam1] = s[jam1]
    loncat = s.index.to_series().diff() != pd.Timedelta(hours=1)
    per_jam[loncat.values & ~jam1] = np.nan
    return per_jam


def terjemahkan(mentah: pd.DataFrame, meta: dict, mode_akumulasi="otomatis"):
    keluar, satuan, jenis, konversi, angin = {}, {}, {}, [], []

    grup = {}
    for kunci in mentah.columns:
        m = meta[kunci]
        grup.setdefault(m["level"], {})[m["nama"]] = kunci

    def tambah(kolom, nilai, sat, jns, asal, ket):
        keluar[kolom] = nilai
        satuan[kolom] = sat
        jenis[kolom] = jns
        konversi.append((kolom, asal, ket))

    for level, var in grup.items():
        akhiran = f" {level}" if level else ""
        terpakai = set()

        for u, v, tinggi in PASANGAN_ANGIN:
            if u in var and v in var:
                U, V = mentah[var[u]], mentah[var[v]]
                dasar = "angin" + (f" {tinggi}" if tinggi else "") + akhiran
                kec = np.sqrt(U ** 2 + V ** 2)
                arah = (270 - np.degrees(np.arctan2(V, U))) % 360
                asal = f"{var[u]}, {var[v]}"
                tambah(f"Kecepatan {dasar} (m/s)", kec, "m/s", "biasa", asal, "akar(u² + v²)")
                angin.append({"nama": dasar, "kec": f"Kecepatan {dasar} (m/s)", "arah": f"Arah {dasar} (°)"})
                tambah(f"Arah {dasar} (°)", arah, "°", "arah", asal,
                       "derajat dari utara searah jarum jam, arah datangnya angin")
                tambah(f"Arah {dasar} (mata angin)", ke_mata_angin(arah), "", "teks", asal,
                       "8 arah mata angin dari kolom derajat")
                terpakai |= {u, v}

        for nama, kunci in var.items():
            if nama in terpakai:
                continue
            s = mentah[kunci]
            m = meta[kunci]
            if nama in KONVERSI:
                tampil, sat, fungsi, ket, jns = KONVERSI[nama]
                if jns in ("jumlah", "fluks"):
                    mode = mode_akumulasi
                    if mode == "otomatis":
                        mode = "sejak_00" if (m["langkah_maks_jam"] or 0) > 12 else "per_jam"
                    if mode == "sejak_00":
                        s = kumulatif_ke_per_jam(s).reindex(mentah.index)
                        ket += "; diubah dari akumulasi sejak 00 UTC ke nilai per jam"
                    else:
                        ket += "; nilai sudah akumulasi 1 jam"
                nilai = fungsi(s) if fungsi else s
                if nama in TIDAK_NEGATIF:
                    nilai = nilai.clip(lower=0)
                kolom = f"{tampil}{akhiran} ({sat})"
                tambah(kolom, nilai, sat, jns, kunci, ket)
                if jns == "arah":
                    tambah(f"{tampil}{akhiran} (mata angin)", ke_mata_angin(nilai), "", "teks",
                           kunci, "8 arah mata angin dari kolom derajat")
            else:
                sat = m["satuan"]
                kolom = f"{m['nama_panjang']}{akhiran}" + (f" ({sat})" if sat else "")
                tambah(kolom, s, sat, "biasa", kunci, "tanpa konversi, nama dan satuan dari file")

        if "t2m" in var and "d2m" in var:
            T = mentah[var["t2m"]] - 273.15
            Td = mentah[var["d2m"]] - 273.15
            rh = 100 * np.exp(17.625 * Td / (243.04 + Td)) / np.exp(17.625 * T / (243.04 + T))
            tambah(f"Kelembapan relatif 2 m{akhiran} (%)", rh.clip(upper=100), "%", "biasa",
                   f"{var['t2m']}, {var['d2m']}", "rumus Magnus dari suhu udara dan titik embun")

    data = pd.DataFrame(keluar, index=mentah.index)
    return data, satuan, jenis, konversi, angin


def proses(daftar_file, lat, lon, zona="UTC", mode_akumulasi="otomatis", lewati_kosong=True, kabar=None):
    """daftar_file: list (path, nama_tampil). kabar(fraksi, teks) untuk progress."""
    semua_seri, meta, grid, peringatan = {}, {}, [], []
    total = len(daftar_file)
    for k, (path, nama) in enumerate(daftar_file):
        sub = (lambda fr, k=k, nama=nama: kabar((k + fr) / total, f"Membaca {nama} ({fr:.0%})")) if kabar else None
        s, m, g, p = baca_titik_grib(path, nama, lat, lon, lewati_kosong, sub)
        for k, v in s.items():
            semua_seri[k] = v if k not in semua_seri else semua_seri[k].combine_first(v)
        meta.update(m)
        grid.extend(g)
        peringatan.extend(p)
    if not semua_seri:
        raise ValueError("Tidak ada variabel yang berhasil diambil dari file.")

    mentah = pd.concat(semua_seri, axis=1).sort_index()
    mentah.index.name = "waktu"
    data, satuan, jenis, konversi, angin = terjemahkan(mentah, meta, mode_akumulasi)

    geser = pd.Timedelta(hours=ZONA_WAKTU[zona])
    data.index = data.index + geser
    data = data.dropna(how="all")

    return Hasil(data=data, mentah=mentah, meta=meta, satuan=satuan, jenis=jenis,
                 konversi=konversi, grid=grid, peringatan=peringatan, zona=zona, angin=angin,
                 lat_input=lat, lon_input=lon, sumber=[n for _, n in daftar_file])


# --------------------------------------------------------------------------
# Grafik deret waktu dan windrose
# --------------------------------------------------------------------------
RESOLUSI = {"Per jam": None, "Per hari": "D", "Per bulan": "MS", "Per tahun": "YS"}

# Kelas kecepatan angin (m/s) mengikuti pembagian umum WRPLOT
TENANG = 0.5
KELAS_ANGIN = [(0.5, 2.1), (2.1, 3.6), (3.6, 5.7), (5.7, 8.8), (8.8, 11.1), (11.1, np.inf)]
LABEL_KELAS = ["0,5-2,1", "2,1-3,6", "3,6-5,7", "5,7-8,8", "8,8-11,1", "≥ 11,1"]
WARNA_KELAS = ["#c6dbef", "#6baed6", "#2171b5", "#fdae6b", "#e6550d", "#a63603"]
SEKTOR = {
    16: ["U", "UTL", "TL", "TTL", "T", "TTG", "TG", "STG", "S", "SBD", "BD", "BBD", "B", "BBL", "BL", "UBL"],
    8: ["U", "TL", "T", "TG", "S", "BD", "B", "BL"],
}
NAMA_BULAN = ["Januari", "Februari", "Maret", "April", "Mei", "Juni", "Juli",
              "Agustus", "September", "Oktober", "November", "Desember"]
MUSIM = {
    "Des-Jan-Feb (DJF)": [12, 1, 2],
    "Mar-Apr-Mei (MAM)": [3, 4, 5],
    "Jun-Jul-Agu (JJA)": [6, 7, 8],
    "Sep-Okt-Nov (SON)": [9, 10, 11],
}


def resolusi_otomatis(index) -> str:
    """Pilih resolusi grafik dari panjang rentang data."""
    if len(index) < 2:
        return "Per jam"
    rentang = index.max() - index.min()
    if rentang <= pd.Timedelta(days=7):
        return "Per jam"
    if rentang <= pd.Timedelta(days=120):
        return "Per hari"
    if rentang <= pd.Timedelta(days=3 * 366):
        return "Per bulan"
    return "Per tahun"


KOLOM_AGREGASI = ["waktu", "label", "nilai", "min", "maks", "p10", "p90", "n", "kelengkapan"]


def agregasi(s: pd.Series, jenis: str, resolusi: str) -> pd.DataFrame:
    """Ringkas deret waktu per periode. 'nilai' = total untuk variabel akumulasi,
    rata-rata untuk lainnya. p10/p90 = persentil 10 dan 90 data di periode itu."""
    s = s.dropna().sort_index()
    frek = RESOLUSI[resolusi]
    if s.empty:
        return pd.DataFrame(columns=KOLOM_AGREGASI)
    if frek is None:
        return pd.DataFrame({"waktu": s.index, "label": s.index.strftime("%Y-%m-%d %H:%M"),
                             "nilai": s.values, "min": s.values, "maks": s.values,
                             "p10": s.values, "p90": s.values, "n": 1, "kelengkapan": 1.0})

    g = s.resample(frek)
    df = pd.DataFrame({
        "nilai": g.sum(min_count=1) if jenis == "jumlah" else g.mean(),
        "min": g.min(), "maks": g.max(), "p10": g.quantile(0.1), "p90": g.quantile(0.9),
        "n": g.count(),
    }).dropna(subset=["nilai"])

    # Kelengkapan: jumlah data dibanding yang seharusnya ada pada periode itu
    langkah = s.index.to_series().diff().median()
    if pd.isna(langkah) or langkah <= pd.Timedelta(0):
        langkah = pd.Timedelta(hours=1)
    ofs = pd.tseries.frequencies.to_offset(frek)
    harusnya = [(t + ofs - t) / langkah for t in df.index]
    df["kelengkapan"] = np.clip(df["n"].values / np.maximum(harusnya, 1), 0, 1)

    fmt = {"D": "%Y-%m-%d", "MS": "%Y-%m", "YS": "%Y"}[frek]
    df["label"] = df.index.strftime(fmt)
    df.index.name = "waktu"
    return df.reset_index()[KOLOM_AGREGASI]


def grafik_deret(df: pd.DataFrame, nama: str, satuan: str, jenis: str, resolusi: str):
    """Grafik Altair: batang untuk total akumulasi, garis rata-rata dengan
    pita persentil 10-90 untuk variabel lain."""
    import altair as alt

    alt.data_transformers.disable_max_rows()
    df = df.copy()
    df["lengkap"] = np.where(df["kelengkapan"] >= 0.9, "Lengkap", "Belum lengkap")
    df["kelengkapan_persen"] = (df["kelengkapan"] * 100).round(0)

    per_jam = RESOLUSI[resolusi] is None
    sumbu_x = (alt.X("waktu:T", title=None) if resolusi in ("Per jam", "Per hari")
               else alt.X("label:O", title=None, sort=None, axis=alt.Axis(labelAngle=0 if resolusi == "Per tahun" else -45, labelOverlap="greedy")))
    judul_y = f"{'Total' if jenis == 'jumlah' else 'Rata-rata'} {nama.split(' (')[0].lower()} ({satuan})" if not per_jam \
        else f"{nama.split(' (')[0]} ({satuan})"
    tip = [alt.Tooltip("label:N", title="Periode"),
           alt.Tooltip("nilai:Q", title="Total" if jenis == "jumlah" else "Rata-rata", format=".2f")]
    if not per_jam:
        tip += [alt.Tooltip("p10:Q", title="Persentil 10", format=".2f"),
                alt.Tooltip("p90:Q", title="Persentil 90", format=".2f"),
                alt.Tooltip("min:Q", title="Minimum", format=".2f"),
                alt.Tooltip("maks:Q", title="Maksimum", format=".2f"),
                alt.Tooltip("n:Q", title="Jumlah data"),
                alt.Tooltip("kelengkapan_persen:Q", title="Kelengkapan (%)")]

    dasar = alt.Chart(df).encode(x=sumbu_x)
    if jenis == "jumlah":
        grafik = dasar.mark_bar(opacity=0.9).encode(
            y=alt.Y("nilai:Q", title=judul_y),
            opacity=alt.Opacity("lengkap:N", scale=alt.Scale(domain=["Lengkap", "Belum lengkap"], range=[0.9, 0.4]),
                                legend=alt.Legend(title=None, orient="top")),
            tooltip=tip)
    elif per_jam:
        grafik = dasar.mark_line(strokeWidth=1).encode(y=alt.Y("nilai:Q", title=judul_y), tooltip=tip)
    else:
        pita = dasar.mark_area(opacity=0.25).encode(y=alt.Y("p10:Q", title=judul_y), y2="p90:Q")
        garis = dasar.mark_line(strokeWidth=2).encode(y="nilai:Q")
        titik = dasar.mark_point(filled=True, size=60).encode(
            y="nilai:Q", tooltip=tip,
            shape=alt.Shape("lengkap:N", scale=alt.Scale(domain=["Lengkap", "Belum lengkap"], range=["circle", "triangle"]),
                            legend=alt.Legend(title=None, orient="top")))
        grafik = pita + garis + titik
    return grafik.properties(height=360)


def saring_periode(index, pilihan: str):
    """Mask boolean untuk filter windrose: semua data, musim, bulan, atau tahun."""
    if pilihan == "Semua data":
        return np.ones(len(index), dtype=bool)
    if pilihan in MUSIM:
        return index.month.isin(MUSIM[pilihan])
    if pilihan in NAMA_BULAN:
        return index.month == NAMA_BULAN.index(pilihan) + 1
    if pilihan.startswith("Tahun "):
        return index.year == int(pilihan.split()[1])
    return np.ones(len(index), dtype=bool)


def tabel_windrose(kec: pd.Series, arah: pd.Series, n_sektor=16):
    """Persentase kejadian per sektor arah x kelas kecepatan.
    Mengembalikan (tabel %, persen tenang, jumlah data)."""
    ok = kec.notna() & arah.notna()
    kec, arah = kec[ok].values, arah[ok].values
    total = len(kec)
    label = SEKTOR[n_sektor]
    tabel = pd.DataFrame(0.0, index=label, columns=LABEL_KELAS)
    if total == 0:
        return tabel, 0.0, 0
    lebar = 360 / n_sektor
    sektor = (np.floor(((arah + lebar / 2) % 360) / lebar).astype(int)) % n_sektor
    for (a, b), nama in zip(KELAS_ANGIN, LABEL_KELAS):
        m = (kec >= a) & (kec < b)
        tabel[nama] = np.bincount(sektor[m], minlength=n_sektor) / total * 100
    tenang = float((kec < TENANG).sum() / total * 100)
    return tabel, tenang, total


def gambar_windrose(tabel: pd.DataFrame, judul=""):
    import plotly.graph_objects as go

    n = len(tabel)
    sudut = [k * 360 / n for k in range(n)]
    fig = go.Figure()
    for nama, warna in zip(tabel.columns, WARNA_KELAS):
        if tabel[nama].sum() <= 0:
            continue
        fig.add_trace(go.Barpolar(
            r=tabel[nama].values, theta=sudut, width=[360 / n * 0.92] * n,
            name=f"{nama} m/s", marker_color=warna, marker_line_width=0,
            customdata=tabel.index,
            hovertemplate="%{customdata}: %{r:.2f}%<extra>" + nama + " m/s</extra>"))
    fig.update_layout(
        title=judul, height=520, margin=dict(t=60, b=30, l=30, r=30),
        legend=dict(title="Kecepatan", orientation="v"),
        polar=dict(
            bargap=0,
            angularaxis=dict(direction="clockwise", rotation=90, tickmode="array",
                             tickvals=sudut, ticktext=list(tabel.index)),
            radialaxis=dict(ticksuffix="%", angle=90, tickangle=90),
        ),
    )
    return fig


# --------------------------------------------------------------------------
# Excel
# --------------------------------------------------------------------------
def buat_excel(h: Hasil, sertakan_mentah=False) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, NamedStyle, PatternFill, Side
    from openpyxl.utils import get_column_letter

    FONT = "Century Gothic"
    garis = Side(style="thin", color="FFD9D9D9")
    KOTAK = Border(left=garis, right=garis, top=garis, bottom=garis)
    KUNING = PatternFill("solid", fgColor="FFFFFF00")

    wb = Workbook()
    gaya = {
        "th": NamedStyle("th", font=Font(name=FONT, size=10, bold=True),
                         alignment=Alignment(horizontal="center", vertical="center", wrap_text=True)),
        "waktu": NamedStyle("waktu", font=Font(name=FONT, size=10), number_format="yyyy-mm-dd hh:mm",
                            alignment=Alignment(horizontal="left", vertical="center")),
        "angka": NamedStyle("angka", font=Font(name=FONT, size=10), number_format="0.000",
                            alignment=Alignment(horizontal="right", vertical="center")),
        "teks": NamedStyle("teks", font=Font(name=FONT, size=10),
                           alignment=Alignment(horizontal="left", vertical="center", indent=1)),
        "angka_mentah": NamedStyle("angka_mentah", font=Font(name=FONT, size=10), number_format="0.0######",
                                   alignment=Alignment(horizontal="right", vertical="center")),
        "th_k": NamedStyle("th_k", font=Font(name=FONT, size=10, bold=True), border=KOTAK,
                           alignment=Alignment(horizontal="center", vertical="center", wrap_text=True)),
        "teks_k": NamedStyle("teks_k", font=Font(name=FONT, size=10), border=KOTAK,
                             alignment=Alignment(horizontal="left", vertical="center", wrap_text=True, indent=1)),
        "angka_k": NamedStyle("angka_k", font=Font(name=FONT, size=10), border=KOTAK, number_format="0.00",
                              alignment=Alignment(horizontal="right", vertical="center")),
        "persen_k": NamedStyle("persen_k", font=Font(name=FONT, size=10), border=KOTAK, number_format="0.0%",
                               alignment=Alignment(horizontal="right", vertical="center")),
        "strip_k": NamedStyle("strip_k", font=Font(name=FONT, size=10), border=KOTAK,
                              alignment=Alignment(horizontal="center", vertical="center")),
        "label": NamedStyle("label", font=Font(name=FONT, size=10, bold=True),
                            alignment=Alignment(horizontal="left", vertical="center")),
    }
    for g in gaya.values():
        wb.add_named_style(g)

    def lebar(teks, minimum=12, maksimum=30):
        return max(minimum, min(maksimum, round(len(str(teks)) * 1.15 / 1.6 + 3)))

    def tulis_tabel_waktu(ws, df, judul_waktu, kolom_teks=(), gaya_angka="angka"):
        ws.append([judul_waktu] + list(df.columns))
        for c in range(1, df.shape[1] + 2):
            ws.cell(1, c).style = "th"
        ws.row_dimensions[1].height = 36
        ws.column_dimensions["A"].width = 20
        for k, nama in enumerate(df.columns, start=2):
            ws.column_dimensions[get_column_letter(k)].width = lebar(nama, 14, 24)
        gaya_kolom = ["waktu"] + ["teks" if c in kolom_teks else gaya_angka for c in df.columns]
        nilai = df.astype(object).where(df.notna(), None)
        for r, (t, baris) in enumerate(zip(df.index, nilai.itertuples(index=False)), start=2):
            ws.append([t.to_pydatetime()] + [None if v is pd.NA else v for v in baris])
            for c, gs in enumerate(gaya_kolom, start=1):
                ws.cell(r, c).style = gs
        ws.sheet_format.defaultRowHeight = 18
        ws.sheet_format.customHeight = True
        ws.auto_filter.ref = f"A1:{get_column_letter(df.shape[1] + 1)}{df.shape[0] + 1}"

    # ---- Sheet Data
    ws = wb.active
    ws.title = "Data"
    kolom_teks = [c for c, j in h.jenis.items() if j == "teks"]
    tulis_tabel_waktu(ws, h.data, f"Waktu ({h.zona.split(' ')[0]})", kolom_teks)
    n_akhir = h.data.shape[0] + 1
    huruf = {nama: get_column_letter(k) for k, nama in enumerate(h.data.columns, start=2)}

    # ---- Sheet Ringkasan (formula, ikut berubah kalau Data diedit)
    wr = wb.create_sheet("Ringkasan")
    judul = ["Variabel", "Satuan", "Jumlah data", "Rata-rata", "Minimum", "Maksimum", "Total"]
    wr.append(judul)
    for c in range(1, len(judul) + 1):
        wr.cell(1, c).style = "th_k"
    wr.row_dimensions[1].height = 22
    r = 2
    for nama in h.data.columns:
        if h.jenis[nama] in ("arah", "teks"):
            continue
        rg = f"Data!{huruf[nama]}2:{huruf[nama]}{n_akhir}"
        wr.cell(r, 1, nama).style = "teks_k"
        wr.cell(r, 2, h.satuan[nama]).style = "teks_k"
        wr.cell(r, 3, f"=COUNT({rg})").style = "angka_k"
        wr.cell(r, 3).number_format = "0"
        wr.cell(r, 4, f'=IF(COUNT({rg})=0,"-",AVERAGE({rg}))').style = "angka_k"
        wr.cell(r, 5, f'=IF(COUNT({rg})=0,"-",MIN({rg}))').style = "angka_k"
        wr.cell(r, 6, f'=IF(COUNT({rg})=0,"-",MAX({rg}))').style = "angka_k"
        if h.jenis[nama] == "jumlah":
            wr.cell(r, 7, f"=SUM({rg})").style = "angka_k"
        else:
            wr.cell(r, 7, "-").style = "strip_k"
        wr.row_dimensions[r].height = 18
        r += 1

    for nama in [c for c in h.data.columns if h.jenis[c] == "teks"]:
        r += 1
        wr.cell(r, 1, "Frekuensi " + nama.replace(" (mata angin)", "")[:1].lower() + nama.replace(" (mata angin)", "")[1:]).style = "label"
        wr.row_dimensions[r].height = 22
        r += 1
        for c, t in enumerate(["Arah", "Jumlah data", "Persentase"], start=1):
            wr.cell(r, c, t).style = "th_k"
        wr.row_dimensions[r].height = 22
        awal = r + 1
        rg = f"Data!{huruf[nama]}2:{huruf[nama]}{n_akhir}"
        for k, arah in enumerate(NAMA_ARAH):
            rr = awal + k
            wr.cell(rr, 1, arah).style = "teks_k"
            wr.cell(rr, 2, f'=COUNTIF({rg},"{arah}")').style = "angka_k"
            wr.cell(rr, 2).number_format = "0"
            wr.cell(rr, 3, f"=IF(SUM($B${awal}:$B${awal + 7})=0,0,B{rr}/SUM($B${awal}:$B${awal + 7}))").style = "persen_k"
            wr.row_dimensions[rr].height = 18
        r = awal + 8
    for kol, w in zip("ABCDEFG", [46, 10, 13, 13, 13, 13, 13]):
        wr.column_dimensions[kol].width = w

    # ---- Sheet Windrose (formula COUNTIFS dari sheet Data)
    ww = None
    if h.angin:
        ww = wb.create_sheet("Windrose")
        n_kls = len(KELAS_ANGIN)
        kol_total = 4 + n_kls                       # kolom "Total"
        r = 1
        for a in h.angin:
            kec = f"Data!${huruf[a['kec']]}$2:${huruf[a['kec']]}${n_akhir}"
            arah = f"Data!${huruf[a['arah']]}$2:${huruf[a['arah']]}${n_akhir}"

            ww.cell(r, 1, f"Windrose {a['nama']}").style = "label"
            ww.row_dimensions[r].height = 22
            r += 1
            ww.cell(r, 1, "Jumlah data").style = "teks"
            sel_n = ww.cell(r, 2, f"=COUNT({kec})")
            sel_n.style = "angka"
            sel_n.number_format = "0"
            ref_n = f"$B${r}"
            ww.row_dimensions[r].height = 18
            r += 2

            # Header: nama kelas, lalu batas bawah dan atas kecepatan (m/s)
            baris_h, baris_lo, baris_hi = r, r + 1, r + 2
            for c, t in enumerate(["Arah", "Dari (°)", "Sampai (°)"], start=1):
                ww.cell(baris_h, c, t).style = "th_k"
                ww.cell(baris_lo, c, "Kecepatan dari (m/s)" if c == 1 else None).style = "teks_k" if c == 1 else "strip_k"
                ww.cell(baris_hi, c, "Kecepatan sampai (m/s)" if c == 1 else None).style = "teks_k" if c == 1 else "strip_k"
            for k, ((lo, hi), lab) in enumerate(zip(KELAS_ANGIN, LABEL_KELAS)):
                c = 4 + k
                ww.cell(baris_h, c, f"{lab} m/s").style = "th_k"
                ww.cell(baris_lo, c, lo).style = "angka_k"
                ww.cell(baris_hi, c, 999 if np.isinf(hi) else hi).style = "angka_k"
                ww.cell(baris_lo, c).number_format = ww.cell(baris_hi, c).number_format = "0.0"
            ww.cell(baris_h, kol_total, "Total").style = "th_k"
            ww.cell(baris_lo, kol_total).style = "strip_k"
            ww.cell(baris_hi, kol_total).style = "strip_k"
            for rr, tinggi in ((baris_h, 30), (baris_lo, 18), (baris_hi, 18)):
                ww.row_dimensions[rr].height = tinggi

            # 16 sektor arah. Sektor utara melewati 360°, jadi dihitung dua bagian.
            awal = baris_hi + 1
            label16 = SEKTOR[16]
            for s_idx, lab in enumerate(label16):
                rr = awal + s_idx
                dari = (s_idx * 22.5 - 11.25) % 360
                sampai = (s_idx * 22.5 + 11.25) % 360
                ww.cell(rr, 1, lab).style = "teks_k"
                ww.cell(rr, 2, dari).style = "angka_k"
                ww.cell(rr, 3, sampai).style = "angka_k"
                for k in range(n_kls):
                    c = 4 + k
                    col = get_column_letter(c)
                    lo, hi = f"{col}${baris_lo}", f"{col}${baris_hi}"
                    syarat_kec = f'{kec},">="&{lo},{kec},"<"&{hi}'
                    f = (f'=IF({ref_n}=0,0,IF($B{rr}<$C{rr},'
                         f'COUNTIFS({arah},">="&$B{rr},{arah},"<"&$C{rr},{syarat_kec}),'
                         f'COUNTIFS({arah},">="&$B{rr},{syarat_kec})+COUNTIFS({arah},"<"&$C{rr},{syarat_kec}))/{ref_n})')
                    ww.cell(rr, c, f).style = "persen_k"
                ww.cell(rr, kol_total,
                        f"=SUM({get_column_letter(4)}{rr}:{get_column_letter(3 + n_kls)}{rr})").style = "persen_k"
                ww.row_dimensions[rr].height = 18

            akhir = awal + len(label16) - 1
            rr = akhir + 1
            ww.cell(rr, 1, "Jumlah per kelas").style = "th_k"
            ww.cell(rr, 2).style = "strip_k"
            ww.cell(rr, 3).style = "strip_k"
            for c in range(4, kol_total + 1):
                col = get_column_letter(c)
                ww.cell(rr, c, f"=SUM({col}{awal}:{col}{akhir})").style = "persen_k"
            ww.row_dimensions[rr].height = 20
            rr += 1
            ww.cell(rr, 1, f"Tenang (< {TENANG:g} m/s)".replace(".", ",")).style = "teks_k"
            ww.cell(rr, 2).style = "strip_k"
            ww.cell(rr, 3).style = "strip_k"
            for c in range(4, kol_total):
                ww.cell(rr, c).style = "strip_k"
            ww.cell(rr, kol_total,
                    f'=IF({ref_n}=0,0,COUNTIFS({kec},"<"&{get_column_letter(4)}${baris_lo})/{ref_n})').style = "persen_k"
            ww.row_dimensions[rr].height = 18
            rr += 1
            ww.cell(rr, 1, "Total").style = "th_k"
            for c in range(2, kol_total):
                ww.cell(rr, c).style = "strip_k"
            ww.cell(rr, kol_total, f"={get_column_letter(kol_total)}{rr - 2}+{get_column_letter(kol_total)}{rr - 1}").style = "persen_k"
            ww.row_dimensions[rr].height = 20
            r = rr + 3

        ww.cell(r - 2, 1, "Persentase dihitung dari seluruh data. Arah = arah datangnya angin, "
                          "batas sektor: dari ≤ arah < sampai.").font = Font(name=FONT, size=8, color="FF7F7F7F")
        ww.column_dimensions["A"].width = 24
        for c in range(2, kol_total + 1):
            ww.column_dimensions[get_column_letter(c)].width = 12

    # ---- Sheet Info
    wi = wb.create_sheet("Info")
    baris_info = [
        ("Dibuat", datetime.now().strftime("%Y-%m-%d %H:%M")),
        ("File sumber", ", ".join(h.sumber)),
        ("Koordinat input", f"{h.lat_input:.4f}, {h.lon_input:.4f}"),
        ("Zona waktu", h.zona),
        ("Periode", f"{h.data.index.min():%Y-%m-%d %H:%M} s.d. {h.data.index.max():%Y-%m-%d %H:%M}"
         if len(h.data) else "-"),
        ("Jumlah baris", len(h.data)),
    ]
    for k, (a, b) in enumerate(baris_info, start=1):
        wi.cell(k, 1, a).style = "label"
        wi.cell(k, 2, b).style = "teks"
        wi.row_dimensions[k].height = 18
    r = len(baris_info) + 2
    for p in h.peringatan:
        sel = wi.cell(r, 1, f"Perhatian: {p}")
        sel.style = "teks"
        sel.fill = KUNING
        r += 1
    r += 0 if not h.peringatan else 1

    wi.cell(r, 1, "Grid terpilih").style = "label"
    r += 1
    judul_grid = ["File", "Variabel", "Catatan", "Lintang grid", "Bujur grid", "Jarak (km)", "Resolusi"]
    for c, t in enumerate(judul_grid, start=1):
        wi.cell(r, c, t).style = "th_k"
    wi.row_dimensions[r].height = 22
    for g in h.grid:
        r += 1
        isi = [g.sumber, g.variabel, g.catatan or "-", g.lat_grid, g.lon_grid, g.jarak_km, g.resolusi]
        for c, v in enumerate(isi, start=1):
            wi.cell(r, c, v).style = "angka_k" if isinstance(v, float) else "teks_k"
        wi.cell(r, 4).number_format = wi.cell(r, 5).number_format = "0.00"
        if g.catatan:
            wi.cell(r, 3).fill = KUNING
        wi.row_dimensions[r].height = 30 if g.catatan else 18

    r += 2
    wi.cell(r, 1, "Asal dan konversi kolom").style = "label"
    r += 1
    for c, t in enumerate(["Kolom di sheet Data", "Variabel asal", "Konversi"], start=1):
        wi.cell(r, c, t).style = "th_k"
    wi.row_dimensions[r].height = 22
    for kol, asal, ket in h.konversi:
        r += 1
        for c, v in enumerate([kol, asal, ket], start=1):
            wi.cell(r, c, v).style = "teks_k"
        wi.row_dimensions[r].height = 30 if (len(kol) > 30 or len(ket) > 42) else 18

    r += 2
    catatan = [
        "Arah angin dan gelombang adalah arah datang, diukur dari utara searah jarum jam.",
        "Nilai akumulasi (hujan, radiasi) dicatat pada akhir jam akumulasinya.",
        "Kolom angka di sheet Data disimpan tanpa pembulatan; tampilan 3 desimal hanya format sel.",
    ]
    for t in catatan:
        wi.cell(r, 1, t).font = Font(name=FONT, size=8, color="FF7F7F7F")
        wi.row_dimensions[r].height = 18
        r += 1
    for kol, w in zip("ABCDEFG", [36, 28, 48, 13, 13, 12, 15]):
        wi.column_dimensions[kol].width = w

    for lembar in [x for x in (wr, wi, ww) if x is not None]:
        lembar.page_setup.orientation = "landscape"
        lembar.page_setup.fitToWidth = 1
        lembar.page_setup.fitToHeight = 0
        lembar.sheet_properties.pageSetUpPr.fitToPage = True

    # ---- Sheet Data mentah (opsional)
    if sertakan_mentah:
        wm = wb.create_sheet("Data mentah")
        mentah = h.mentah.copy()
        mentah.columns = [
            f"{k} ({h.meta[k]['satuan']})" if h.meta[k]["satuan"] else k for k in mentah.columns
        ]
        tulis_tabel_waktu(wm, mentah, "Waktu (UTC)", gaya_angka="angka_mentah")

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# --------------------------------------------------------------------------
# Antarmuka Streamlit
# --------------------------------------------------------------------------
def main():
    import streamlit as st

    st.set_page_config(page_title="GRIB ke Excel", layout="wide")
    st.title("GRIB ke Excel")
    st.caption("Ambil deret waktu ERA5 dari file GRIB pada titik grid terdekat, lalu simpan sebagai Excel.")

    with st.sidebar:
        st.header("1. File GRIB")
        cara = st.radio("Sumber file", ["Unggah", "Path di komputer"], horizontal=True,
                        help="File besar lebih cepat dibaca langsung dari path tanpa diunggah.")
        unggahan, path_teks = [], ""
        if cara == "Unggah":
            unggahan = st.file_uploader("Pilih file .grib", accept_multiple_files=True,
                                        type=["grib", "grb", "grib1", "grib2", "grb2"])
        else:
            path_teks = st.text_area("Path file, satu per baris",
                                     placeholder="D:/data/era5_angin_2024.grib")

        st.header("2. Lokasi")
        c1, c2 = st.columns(2)
        lat = c1.number_input("Lintang", min_value=-90.0, max_value=90.0, value=None,
                              format="%.4f", placeholder="-7.2500", help="Negatif untuk lintang selatan")
        lon = c2.number_input("Bujur", min_value=-180.0, max_value=360.0, value=None,
                              format="%.4f", placeholder="112.7500", help="Negatif untuk bujur barat")

        st.header("3. Pengaturan")
        zona = st.selectbox("Zona waktu di Excel", list(ZONA_WAKTU), index=1)
        mode_label = {
            "Otomatis": "otomatis",
            "Per jam (ERA5)": "per_jam",
            "Akumulasi sejak 00 UTC (ERA5-Land)": "sejak_00",
        }
        mode = st.selectbox("Variabel akumulasi (hujan, radiasi)", list(mode_label),
                            help="ERA5 menyimpan hujan per jam. ERA5-Land menyimpan total sejak 00 UTC, "
                                 "jadi perlu diselisihkan. Mode otomatis menebak dari panjang step di file.")
        lewati = st.checkbox("Lewati grid kosong", value=True,
                             help="Contoh: data gelombang tidak ada di daratan. Kalau grid terdekat kosong, "
                                  "dipakai grid berisi yang paling dekat.")
        mentah = st.checkbox("Sertakan sheet data mentah", value=False)
        jalan = st.button("Proses", type="primary", width="stretch")

    if jalan:
        if lat is None or lon is None:
            st.warning("Isi lintang dan bujur dulu.")
            st.stop()
        sementara, daftar = [], []
        try:
            if cara == "Unggah":
                if not unggahan:
                    st.warning("Belum ada file yang diunggah.")
                    st.stop()
                for f in unggahan:
                    tmp = tempfile.NamedTemporaryFile(delete=False, suffix=".grib")
                    tmp.write(f.getbuffer())
                    tmp.close()
                    sementara.append(tmp.name)
                    daftar.append((tmp.name, f.name))
            else:
                for p in [x.strip().strip('"') for x in path_teks.splitlines() if x.strip()]:
                    if not os.path.isfile(p):
                        st.error(f"File tidak ditemukan: {p}")
                        st.stop()
                    daftar.append((p, os.path.basename(p)))
                if not daftar:
                    st.warning("Isi path file dulu.")
                    st.stop()

            bar = st.progress(0.0, text="Mulai membaca GRIB...")
            hasil = proses(daftar, float(lat), float(lon), zona, mode_label[mode], lewati,
                           kabar=lambda fr, teks: bar.progress(min(fr, 1.0) * 0.9, text=teks))
            bar.progress(0.92, text="Menulis Excel...")
            xlsx = buat_excel(hasil, mentah)
            bar.empty()
            st.session_state["hasil"] = hasil
            st.session_state["xlsx"] = xlsx
        except Exception as e:  # noqa: BLE001
            pesan = str(e)
            st.error(f"Gagal memproses file: {pesan}")
            if "eccodes" in pesan.lower() or "ecCodes" in pesan:
                st.info("Library ecCodes belum terpasang. Coba: pip install eccodes")
            st.stop()
        finally:
            for p in sementara:
                try:
                    os.remove(p)
                except OSError:
                    pass

    hasil: Hasil | None = st.session_state.get("hasil")
    if hasil is None:
        st.info("Pilih file GRIB, isi koordinat, lalu tekan Proses.")
        return

    for p in hasil.peringatan:
        st.warning(p)

    g0 = hasil.grid[0]
    m1, m2, m3, m4 = st.columns(4)
    m1.metric("Grid terpilih", f"{g0.lat_grid:.2f}, {g0.lon_grid:.2f}")
    m2.metric("Jarak dari input", f"{g0.jarak_km:.1f} km")
    m3.metric("Jumlah baris", f"{len(hasil.data):,}".replace(",", "."))
    m4.metric("Kolom hasil", hasil.data.shape[1])

    st.subheader("Grid per dataset")
    st.dataframe(pd.DataFrame([g.__dict__ for g in hasil.grid]).rename(columns={
        "sumber": "File", "variabel": "Variabel", "lat_grid": "Lintang grid", "lon_grid": "Bujur grid",
        "jarak_km": "Jarak (km)", "resolusi": "Resolusi", "catatan": "Catatan"}),
        hide_index=True, width="stretch")

    peta = pd.DataFrame(
        [{"lat": hasil.lat_input, "lon": hasil.lon_input if hasil.lon_input <= 180 else hasil.lon_input - 360,
          "warna": "#d62728"}]
        + [{"lat": g.lat_grid, "lon": g.lon_grid, "warna": "#1f77b4"} for g in hasil.grid])
    st.map(peta, latitude="lat", longitude="lon", color="warna", size=600, zoom=8)
    st.caption("Merah: koordinat input. Biru: grid yang dipakai.")

    nama_file = f"era5_{hasil.lat_input:.3f}_{hasil.lon_input:.3f}.xlsx".replace("-", "m")

    # ---------------- Grafik deret waktu
    angka = [c for c, j in hasil.jenis.items() if j not in ("teks", "arah")]
    if angka:
        st.subheader("Grafik deret waktu")
        auto = resolusi_otomatis(hasil.data.index)
        c1, c2 = st.columns([2, 1])
        nama = c1.selectbox("Variabel", angka)
        opsi = [f"Otomatis ({auto.lower()})"] + list(RESOLUSI)
        pilih_res = c2.selectbox("Resolusi", opsi)
        res = auto if pilih_res.startswith("Otomatis") else pilih_res

        jenis = hasil.jenis[nama]
        df = agregasi(hasil.data[nama], jenis, res)
        if res == "Per jam" and len(df) > 20000:
            st.caption(f"{len(df):,} titik per jam, grafik mungkin berat di browser.".replace(",", "."))
        st.altair_chart(grafik_deret(df, nama, hasil.satuan[nama], jenis, res), width="stretch")

        ket = []
        if res != "Per jam":
            ket.append("Garis: total per periode." if jenis == "jumlah"
                       else "Garis: rata-rata per periode. Pita: rentang persentil 10-90 (80% data berada di dalamnya).")
            kurang = df.loc[df["kelengkapan"] < 0.9, "label"].tolist()
            if kurang:
                daftar = ", ".join(kurang[:6]) + (f", dan {len(kurang) - 6} lainnya" if len(kurang) > 6 else "")
                ket.append(f"Periode dengan data belum lengkap (ditandai segitiga/batang pucat): {daftar}.")
        if ket:
            st.caption(" ".join(ket))

    # ---------------- Windrose
    if hasil.angin:
        st.subheader("Windrose")
        idx = hasil.data.index
        opsi_periode = (["Semua data"] + list(MUSIM) + NAMA_BULAN
                        + [f"Tahun {t}" for t in sorted(set(idx.year))])
        c1, c2, c3 = st.columns([2, 2, 1])
        if len(hasil.angin) > 1:
            pilih_angin = c1.selectbox("Ketinggian/level", [a["nama"] for a in hasil.angin])
        else:
            pilih_angin = hasil.angin[0]["nama"]
            c1.text_input("Ketinggian/level", pilih_angin, disabled=True)
        a = next(x for x in hasil.angin if x["nama"] == pilih_angin)
        periode = c2.selectbox("Periode", opsi_periode,
                               help="Musim mengikuti pembagian DJF/MAM/JJA/SON. Bulan dan tahun memakai zona waktu yang dipilih.")
        n_sektor = c3.radio("Sektor", [16, 8], horizontal=True)

        mask = saring_periode(idx, periode)
        tabel, tenang, n = tabel_windrose(hasil.data.loc[mask, a["kec"]], hasil.data.loc[mask, a["arah"]], n_sektor)
        if n == 0:
            st.info("Tidak ada data angin pada periode ini.")
        else:
            judul = f"Windrose {a['nama']}, {periode.lower() if periode == 'Semua data' else periode}"
            st.plotly_chart(gambar_windrose(tabel, judul), width="stretch")
            per_arah = tabel.sum(axis=1)
            kec_rata = hasil.data.loc[mask, a["kec"]].mean()
            m1, m2, m3, m4 = st.columns(4)
            m1.metric("Arah dominan", per_arah.idxmax(), f"{per_arah.max():.1f}% kejadian", delta_color="off")
            m2.metric("Kecepatan rata-rata", f"{kec_rata:.2f} m/s")
            m3.metric("Angin tenang (< 0,5 m/s)", f"{tenang:.1f}%")
            m4.metric("Jumlah data", f"{n:,}".replace(",", "."))
            st.caption("Arah menunjukkan dari mana angin datang. Panjang batang = persentase kejadian dari seluruh "
                       "data di periode ini; angin tenang tidak punya arah sehingga tidak masuk batang. "
                       "Gambar bisa disimpan lewat ikon kamera di pojok kanan atas grafik. "
                       "Tabel windrose lengkap ada di sheet Windrose pada file Excel.")

    st.subheader("Pratinjau data")
    st.dataframe(hasil.data.head(500), width="stretch")

    st.download_button("Unduh Excel", st.session_state["xlsx"], file_name=nama_file,
                       mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                       type="primary")


if __name__ == "__main__":
    main()
